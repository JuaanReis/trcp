import csv
import json
import sqlite3
from datetime import datetime
from pathlib import Path
import questionary
from questionary import Choice
from core.database.leads import list_leads

FIELDS = [
    "name",
    "category",
    "location",
    "website",
    "phone",
    "social_profiles",
    "score",
    "status",
    "created_at",
]

def _serialize(lead) -> dict:
    return {
        "name": lead.name,
        "category": lead.category,
        "location": lead.location,
        "website": lead.website,
        "phone": lead.phone,
        "social_profiles": lead.social_profiles,
        "score": lead.score,
        "status": lead.status,
        "created_at": lead.created_at.isoformat(),
    }


def _export_csv(path: Path, rows: list[dict]):
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=FIELDS,
            delimiter=";",
        )
        writer.writeheader()

        for row in rows:
            csv_row = {
                **row,
                "social_profiles": json.dumps(
                    row["social_profiles"],
                    ensure_ascii=False,
                ),
            }

            # Protege textos externos contra interpretação como fórmulas.
            for key, value in csv_row.items():
                if isinstance(value, str):
                    if value.lstrip().startswith(("=", "+", "-", "@")):
                        csv_row[key] = "'" + value
                    elif value.startswith(("\t", "\r", "\n")):
                        csv_row[key] = "'" + value

            writer.writerow(csv_row)


def _export_json(path: Path, rows: list[dict]):
    with path.open("w", encoding="utf-8") as file:
        json.dump(
            rows,
            file,
            ensure_ascii=False,
            indent=2,
        )
        file.write("\n")


def _export_xlsx(path: Path, rows: list[dict]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()

    try:
        sheet = workbook.active
        sheet.title = "Leads"
        sheet.append(FIELDS)

        for row in rows:
            values = [
                json.dumps(row[field], ensure_ascii=False)
                if isinstance(row[field], (list, dict))
                else row[field]
                for field in FIELDS
            ]

            sheet.append(values)

            # Preserva strings como texto, inclusive as iniciadas por "=".
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="001B3D",
        )

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")

        sheet.row_dimensions[1].height = 24
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for column_index in range(1, len(FIELDS) + 1):
            longest = max(
                len(str(sheet.cell(row_index, column_index).value or ""))
                for row_index in range(1, sheet.max_row + 1)
            )

            letter = get_column_letter(column_index)
            sheet.column_dimensions[letter].width = min(
                max(longest + 2, 16),
                55,
            )

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        workbook.save(path)
    finally:
        workbook.close()


def _export_docx(path: Path, rows: list[dict]):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    document = Document()

    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    document.add_heading("TRCP — Leads", level=0)
    document.add_paragraph(
        f"Generated: {datetime.now():%d/%m/%Y %H:%M}\n"
        f"Total records: {len(rows)}"
    )

    labels = {
        "category": "Category",
        "location": "Location",
        "website": "Website",
        "phone": "Phone",
        "social_profiles": "Social profiles",
        "score": "Score",
        "status": "Status",
        "created_at": "Created",
    }

    for index, row in enumerate(rows, start=1):
        heading = document.add_heading(
            f"{index}. {row['name']}",
            level=2,
        )
        heading.paragraph_format.keep_with_next = True

        for run in heading.runs:
            run.font.color.rgb = RGBColor.from_string("001B3D")

        for field, label in labels.items():
            value = row[field]

            if isinstance(value, (list, dict)):
                value = (
                    json.dumps(value, ensure_ascii=False)
                    if value else "Not informed"
                )
            elif value is None or value == "":
                value = "Not informed"

            paragraph = document.add_paragraph()
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(str(value))

    document.save(path)


def export_leads():
    print("\n[ Export leads ]\n")

    try:
        leads = list_leads()

        if not leads:
            print("No leads found yet.")
            return

        file_format = questionary.select(
            f"Export {len(leads)} saved leads as:",
            choices=[
                Choice("CSV (.csv)", value="csv"),
                Choice("JSON (.json)", value="json"),
                Choice("Excel (.xlsx)", value="xlsx"),
                Choice("Word (.docx)", value="docx"),
                Choice("Back to main menu", value="back"),
            ],
        ).ask()

        if file_format is None or file_format == "back":
            return

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        default_path = (
            Path.cwd()
            / "exports"
            / f"leads_{timestamp}.{file_format}"
        )

        destination = questionary.text(
            "Destination file:",
            default=str(default_path),
            validate=lambda value: (
                bool(value.strip()) or "Enter a destination."
            ),
        ).ask()

        if destination is None:
            return

        path = Path(destination.strip()).expanduser()

        if path.suffix.lower() != f".{file_format}":
            print(f"\nThe filename must end with .{file_format}")
            return

        if path.exists():
            if not path.is_file():
                print("\nThe destination is not a regular file.")
                return

            overwrite = questionary.confirm(
                "File already exists. Overwrite?",
                default=False,
            ).ask()

            if not overwrite:
                print("\nExport canceled.")
                return

        rows = [_serialize(lead) for lead in leads]

        exporters = {
            "csv": _export_csv,
            "json": _export_json,
            "xlsx": _export_xlsx,
            "docx": _export_docx,
        }

        path.parent.mkdir(parents=True, exist_ok=True)
        exporters[file_format](path, rows)

        print(f"\nExported: {len(rows)} leads")
        print(f"File: {path.resolve()}")

    except ImportError as exc:
        print(f"\nMissing export dependency: {exc}")
        print("Install: py -m pip install openpyxl python-docx")

    except (sqlite3.Error, OSError, ValueError, TypeError) as exc:
        print(f"\nExport failed: {exc}")

    finally:
        input("\nPress ENTER to continue...")